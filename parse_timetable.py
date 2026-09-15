#!/usr/bin/env python3
"""
IST Timetable Parser
Reads the IST timetable .xlsx and writes timetable.json.

Key idea: IST classes do not sit in fixed slots. Each class is a merged
block of cells on a 5-minute grid (column D = 08:30). So we read the
merged ranges to get each class's real start and end time.

Usage: python parse_timetable.py timetable.xlsx [timetable.json]
"""
import sys, json, re
from datetime import datetime, timezone
import openpyxl
from openpyxl.utils import get_column_letter

DAYS = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday"]
FIRST_COL = 4            # column D = 08:30
DAY_START = 8 * 60 + 30  # minutes
DAY_END = 16 * 60        # last header is 16:00
MINUTES_PER_COL = 5

def col_to_min(col):
    return DAY_START + (col - FIRST_COL) * MINUTES_PER_COL

def fmt(m):
    return f"{m // 60:02d}:{m % 60:02d}"

def clean(text):
    return " ".join(str(text).split())

ROOM_PATTERNS = [
    re.compile(r"IT\s*Lab\s*#?\s*(\d+)", re.I),
    re.compile(r"Room\s*#?\s*(\d{4})", re.I),
]

def split_subject_room(text):
    text = clean(text)
    for pat in ROOM_PATTERNS:
        m = pat.search(text)
        if m:
            room = f"IT Lab {m.group(1)}" if "lab" in pat.pattern.lower() else m.group(1)
            subject = text[:m.start()].strip(" -#")
            return subject or text, room
    return text, None

def kind_of(subject):
    s = subject.lower()
    if "activit" in s: return "activity"
    if "advisory" in s or "tut" in s: return "advisory"
    if "lab" in s: return "lab"
    return "class"

def find_day_rows(ws):
    rows = {}
    for r in range(1, 15):
        v = ws.cell(r, 1).value
        if isinstance(v, str) and v.strip() in DAYS:
            rows[v.strip()] = r
    return rows

def read_blocks(ws):
    """Return {day: [(start_min, end_min, raw_text), ...]} for one sheet."""
    merged = {}
    for rng in ws.merged_cells.ranges:
        merged[(rng.min_row, rng.min_col)] = rng.max_col
    out = {}
    for day, r in find_day_rows(ws).items():
        items = []
        for c in range(FIRST_COL, ws.max_column + 1):
            v = ws.cell(r, c).value
            if v is None or not str(v).strip():
                continue
            last = merged.get((r, c))
            if last is None:
                # Unmerged cell: assume the block runs until the next filled cell.
                last = c
                while last + 1 <= ws.max_column and ws.cell(r, last + 1).value is None \
                        and (r, last + 1) not in merged:
                    last += 1
                    if last - c >= 23:  # cap at 2 hours
                        break
            start = col_to_min(c)
            end = min(col_to_min(last + 1), DAY_END)
            items.append((start, end, clean(v)))
        out[day] = items
    return out

def sheet_type(name):
    n = name.strip()
    if n.lower().startswith("it lab"): return "lab"
    if n.isdigit(): return "room"
    return "section"

def main(src, dst):
    wb = openpyxl.load_workbook(src)
    data = {
        "meta": {
            "generatedAt": datetime.now(timezone.utc).isoformat(timespec="minutes"),
            "source": src.split("/")[-1],
            "days": DAYS,
            "dayStart": fmt(DAY_START),
            "dayEnd": fmt(DAY_END),
        },
        "sections": {},   # section -> day -> list of classes
        "rooms": {},      # room -> day -> list of busy intervals
        "labs": {},
    }
    warnings = []

    for ws in wb.worksheets:
        name = ws.title.strip()
        kind = sheet_type(name)
        blocks = read_blocks(ws)
        if not blocks:
            warnings.append(f"{name}: no day rows found")
        if kind == "section":
            sched = {}
            for day in DAYS:
                sched[day] = []
                for s, e, raw in blocks.get(day, []):
                    subject, room = split_subject_room(raw)
                    ktype = kind_of(subject)
                    if ktype == "class" and room and room.startswith("IT Lab"):
                        ktype = "lab"
                    sched[day].append({
                        "start": fmt(s), "end": fmt(e),
                        "subject": subject, "room": room,
                        "type": ktype,
                    })
            data["sections"][name] = sched
        else:
            target = data["rooms"] if kind == "room" else data["labs"]
            target[name] = {
                day: [{"start": fmt(s), "end": fmt(e),
                       "subject": split_subject_room(raw)[0]}
                      for s, e, raw in blocks.get(day, [])]
                for day in DAYS
            }

    # Safety net: if a class sheet puts a class in a room but the room's own
    # sheet does not show it, add it, so the room is never shown as free.
    def to_min(t):
        h, m = map(int, t.split(":"))
        return h * 60 + m
    for sec, days in data["sections"].items():
        for day, items in days.items():
            for x in items:
                room = x["room"]
                table = data["rooms"] if room in data["rooms"] else \
                        data["labs"] if room in data["labs"] else None
                if table is None:
                    if room:
                        warnings.append(f"{sec} {day} {x['start']}: unknown room '{room}'")
                    continue
                occ = table[room][day]
                s0, e0 = to_min(x["start"]), to_min(x["end"])
                if not any(to_min(o["start"]) < e0 and s0 < to_min(o["end"]) for o in occ):
                    occ.append({"start": x["start"], "end": x["end"],
                                "subject": x["subject"]})
                    occ.sort(key=lambda o: o["start"])
                    warnings.append(f"{room} sheet missing {sec} {day} "
                                    f"{x['start']}-{x['end']} ({x['subject']})")

    with open(dst, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=1)
    print(f"Wrote {dst}: {len(data['sections'])} sections, "
          f"{len(data['rooms'])} rooms, {len(data['labs'])} labs")
    for w in warnings:
        print("WARNING:", w)

if __name__ == "__main__":
    if len(sys.argv) < 2:
        sys.exit("Usage: python parse_timetable.py timetable.xlsx [timetable.json]")
    main(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else "timetable.json")
